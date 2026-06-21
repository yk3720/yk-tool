"""v0.3 マスターシート読み取り（構成の XLOOKUP 未評価時のフォールバック）。"""

from __future__ import annotations

from openpyxl import Workbook


def _cell_str(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _parse_int_cell(value: object) -> int | None:
    text = _cell_str(value)
    if not text:
        return None
    return int(float(text))


def read_v03_masters(
    workbook: Workbook,
) -> tuple[str, str, dict[int, str], dict[int, str]]:
    """装置名・ユニット・モジュールマスタを読む。"""
    if "装置名" not in workbook.sheetnames:
        raise ValueError("シート「装置名」がありません（v0.3 マスター）")

    ws_device = workbook["装置名"]
    device_rows = list(ws_device.iter_rows(min_row=2, values_only=True))
    if not device_rows:
        raise ValueError("シート「装置名」にデータ行がありません")

    internal_code = _cell_str(device_rows[0][0])
    display_name = _cell_str(device_rows[0][1])
    if not internal_code or not display_name:
        raise ValueError("シート「装置名」2行目に装置製番・装置名を入力してください")

    units: dict[int, str] = {}
    if "ユニット" in workbook.sheetnames:
        for row in workbook["ユニット"].iter_rows(min_row=2, values_only=True):
            uid = _parse_int_cell(row[0] if row else None)
            label = _cell_str(row[1] if row and len(row) > 1 else None)
            if uid is not None and label:
                units[uid] = label

    modules: dict[int, str] = {}
    if "モジュール" in workbook.sheetnames:
        for row in workbook["モジュール"].iter_rows(min_row=2, values_only=True):
            mid = _parse_int_cell(row[0] if row else None)
            label = _cell_str(row[1] if row and len(row) > 1 else None)
            if mid is not None and label:
                modules[mid] = label

    return internal_code, display_name, units, modules
