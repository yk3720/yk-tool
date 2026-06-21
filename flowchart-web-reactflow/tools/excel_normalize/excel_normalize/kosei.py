from __future__ import annotations

from dataclasses import dataclass
from typing import Literal

from openpyxl import Workbook

from .constants import KOSEI_HEADERS, KOSEI_HEADERS_V03, KOSEI_SHEET
from .master_sheets import read_v03_masters


@dataclass(frozen=True)
class KoseiRow:
    internal_code: str
    display_name: str
    unit_label: str
    module_label: str
    sort_index: int
    uin_id: int | None = None
    module_id: int | None = None


@dataclass
class KoseiSheet:
    rows: list[KoseiRow]
    display_name: str
    internal_code: str
    format_version: Literal["v02", "v03"] = "v02"

    @property
    def unit_labels(self) -> list[str]:
        seen: set[str] = set()
        out: list[str] = []
        for row in self.rows:
            if row.unit_label not in seen:
                seen.add(row.unit_label)
                out.append(row.unit_label)
        return out

    def modules_for_unit(self, unit_label: str) -> list[str]:
        return [r.module_label for r in self.rows if r.unit_label == unit_label]

    def unit_sheet_title(self, unit_label: str) -> str:
        if self.format_version == "v03":
            uin_id = self._uin_id_for_unit(unit_label)
            if uin_id is None:
                raise ValueError(f"ユニット「{unit_label}」の UinID が構成にありません")
            return f"U{uin_id}"
        return unit_label

    def _uin_id_for_unit(self, unit_label: str) -> int | None:
        for row in self.rows:
            if row.unit_label == unit_label and row.uin_id is not None:
                return row.uin_id
        return None


def _cell_str(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _parse_int(value: object, *, field: str, row_num: int) -> int:
    text = _cell_str(value)
    if not text:
        raise ValueError(f"シート「{KOSEI_SHEET}」{row_num}行目: {field} が空です")
    try:
        return int(float(text))
    except ValueError as exc:
        raise ValueError(
            f"シート「{KOSEI_SHEET}」{row_num}行目: {field} は整数である必要があります（実際: {text}）"
        ) from exc


def parse_kosei_sheet(workbook: Workbook) -> KoseiSheet:
    if KOSEI_SHEET not in workbook.sheetnames:
        raise ValueError(f"シート「{KOSEI_SHEET}」がありません")

    ws = workbook[KOSEI_SHEET]
    matrix: list[list[str]] = []
    for row in ws.iter_rows(values_only=True):
        matrix.append([_cell_str(v) for v in row])

    if not matrix:
        raise ValueError(f"シート「{KOSEI_SHEET}」が空です")

    header4 = tuple(matrix[0][:4])
    header6 = tuple((matrix[0] + ["", "", ""])[:6])

    if header4 == KOSEI_HEADERS:
        return _parse_kosei_v02(matrix)
    if header6 == KOSEI_HEADERS_V03:
        return _parse_kosei_v03(workbook, matrix)

    raise ValueError(
        f"シート「{KOSEI_SHEET}」1行目が未対応です。"
        f" v0.2: {KOSEI_HEADERS} · v0.3: {KOSEI_HEADERS_V03}（実際: {tuple(matrix[0][:6])}）"
    )


def _finalize_kosei(
    rows: list[KoseiRow], format_version: Literal["v02", "v03"]
) -> KoseiSheet:
    if not rows:
        raise ValueError(f"シート「{KOSEI_SHEET}」にデータ行がありません")

    codes = {r.internal_code for r in rows}
    display_names = {r.display_name for r in rows}
    if len(codes) != 1:
        raise ValueError(
            f"装置製番（社内番号）がファイル内で一意ではありません: {sorted(codes)}"
        )
    if len(display_names) != 1:
        raise ValueError(
            f"装置名がファイル内で一意ではありません: {sorted(display_names)}"
        )

    return KoseiSheet(
        rows=rows,
        internal_code=rows[0].internal_code,
        display_name=rows[0].display_name,
        format_version=format_version,
    )


def _parse_kosei_v02(matrix: list[list[str]]) -> KoseiSheet:
    rows: list[KoseiRow] = []
    for idx, raw in enumerate(matrix[1:], start=0):
        cells = (raw + ["", "", "", ""])[:4]
        if not any(cells):
            continue
        internal_code, display_name, unit_label, module_label = cells
        if not all([internal_code, display_name, unit_label, module_label]):
            raise ValueError(
                f"シート「{KOSEI_SHEET}」{idx + 2}行目: 4列すべて入力してください"
            )
        rows.append(
            KoseiRow(
                internal_code=internal_code,
                display_name=display_name,
                unit_label=unit_label,
                module_label=module_label,
                sort_index=idx,
            )
        )
    return _finalize_kosei(rows, "v02")


def _parse_int_optional(value: object, *, field: str, row_num: int) -> int | None:
    text = _cell_str(value)
    if not text:
        return None
    try:
        return int(float(text))
    except ValueError as exc:
        raise ValueError(
            f"シート「{KOSEI_SHEET}」{row_num}行目: {field} は整数である必要があります（実際: {text}）"
        ) from exc


def _parse_kosei_v03(workbook: Workbook, matrix: list[list[str]]) -> KoseiSheet:
    default_code, default_name, units_map, modules_map = read_v03_masters(workbook)
    ws = workbook[KOSEI_SHEET]
    data_row_count = max(ws.max_row - 1, len(matrix) - 1)
    rows: list[KoseiRow] = []
    for idx in range(data_row_count):
        raw = matrix[idx + 1] if idx + 1 < len(matrix) else []
        cells = (raw + ["", "", "", "", "", ""])[:6]
        row_num = idx + 2
        internal_code, display_name, uin_raw, unit_label, mid_raw, module_label = cells

        uin_id = _parse_int_optional(uin_raw, field="UinID", row_num=row_num)
        if uin_id is None:
            uin_id = idx // 10

        module_id = _parse_int_optional(mid_raw, field="MID", row_num=row_num)
        if module_id is None:
            module_id = uin_id * 100 + (idx % 10)

        if not internal_code:
            internal_code = default_code
        if not display_name:
            display_name = default_name
        if not unit_label:
            unit_label = units_map.get(uin_id, "")
        if not module_label:
            module_label = modules_map.get(module_id, "")

        if not all([internal_code, display_name, unit_label, module_label]):
            raise ValueError(
                f"シート「{KOSEI_SHEET}」{row_num}行目: "
                f"装置製番・装置名・ユニット・モジュールを解決できませんでした"
            )
        rows.append(
            KoseiRow(
                internal_code=internal_code,
                display_name=display_name,
                unit_label=unit_label,
                module_label=module_label,
                sort_index=idx,
                uin_id=uin_id,
                module_id=module_id,
            )
        )
    return _finalize_kosei(rows, "v03")
