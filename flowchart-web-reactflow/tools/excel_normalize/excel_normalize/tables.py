from __future__ import annotations

from dataclasses import dataclass

from openpyxl.utils import range_boundaries
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from .constants import FLOW_COLUMN_COUNT, FLOW_HEADERS, RESERVED_SHEET_NAMES
from .kosei import KoseiSheet


@dataclass(frozen=True)
class FlowTableBlock:
    table_name: str
    module_label: str
    rows: list[list[str]]


def _cell_str(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _unit_short_label(unit_label: str) -> str:
    if unit_label.endswith("ユニット"):
        return unit_label[: -len("ユニット")]
    return unit_label


def resolve_table_module_label(
    table_name: str,
    unit_label: str,
    expected_modules: list[str],
) -> str | None:
    if table_name in expected_modules:
        return table_name

    short = _unit_short_label(unit_label)
    prefixed = f"{short}_{table_name}"
    if prefixed in expected_modules:
        return table_name

    for mod in expected_modules:
        if table_name == f"{short}_{mod}":
            return mod
        if table_name == f"{unit_label}_{mod}":
            return mod

    return None


def _is_header_row(row: list[str]) -> bool:
    if len(row) < FLOW_COLUMN_COUNT:
        return False
    head = [c.strip() for c in row[:FLOW_COLUMN_COUNT]]
    if head == list(FLOW_HEADERS):
        return True
    first = head[0].lower()
    return first in ("id", "ｉｄ")


def _read_table_rows(ws: Worksheet, ref: str) -> list[list[str]]:
    min_col, min_row, max_col, max_row = range_boundaries(ref)
    matrix: list[list[str]] = []
    for r in range(min_row, max_row + 1):
        row: list[str] = []
        for c in range(min_col, min_col + FLOW_COLUMN_COUNT):
            row.append(_cell_str(ws.cell(r, c).value))
        matrix.append(row)
    return matrix


def _data_rows_from_matrix(matrix: list[list[str]]) -> list[list[str]]:
    if not matrix:
        return []
    start = 1 if _is_header_row(matrix[0]) else 0
    data: list[list[str]] = []
    for row in matrix[start:]:
        if not any(cell for cell in row):
            continue
        if len(row) < FLOW_COLUMN_COUNT:
            raise ValueError(
                f"10列必要ですが {len(row)} 列です（先頭セル: {row[0] if row else '空'}）"
            )
        normalized = row[:FLOW_COLUMN_COUNT]
        while len(normalized) < FLOW_COLUMN_COUNT:
            normalized.append("")
        data.append(normalized)
    return data


def extract_unit_sheet_tables(
    ws: Worksheet,
    unit_label: str,
    expected_modules: list[str],
) -> list[FlowTableBlock]:
    if not ws.tables:
        raise ValueError(
            f"シート「{ws.title}」: Excel テーブルがありません。"
            f"各動作を「挿入 → テーブル」で登録してください（期待動作: {', '.join(expected_modules)}）"
        )

    blocks: list[FlowTableBlock] = []
    seen_modules: set[str] = set()

    for table in ws.tables.values():
        table_name = table.name
        matrix = _read_table_rows(ws, table.ref)
        try:
            data_rows = _data_rows_from_matrix(matrix)
        except ValueError as exc:
            raise ValueError(
                f"シート「{ws.title}」· テーブル「{table_name}」: {exc}"
            ) from exc

        if not data_rows:
            raise ValueError(
                f"シート「{ws.title}」· テーブル「{table_name}」: データ行がありません"
            )

        module_label = resolve_table_module_label(
            table_name, unit_label, expected_modules
        )
        if module_label is None:
            raise ValueError(
                f"シート「{ws.title}」· テーブル「{table_name}」: "
                f"構成シートの動作名と一致しません（期待: {', '.join(expected_modules)}）"
            )
        if module_label in seen_modules:
            raise ValueError(
                f"シート「{ws.title}」: 動作「{module_label}」に対応するテーブルが重複しています"
            )
        seen_modules.add(module_label)
        blocks.append(
            FlowTableBlock(
                table_name=table_name,
                module_label=module_label,
                rows=data_rows,
            )
        )

    return blocks


def list_unit_sheets(workbook: Workbook, kosei: KoseiSheet) -> dict[str, Worksheet]:
    sheets: dict[str, Worksheet] = {}
    for name in workbook.sheetnames:
        if name in RESERVED_SHEET_NAMES:
            continue
        if name.startswith("_"):
            continue
        sheets[name] = workbook[name]

    unit_sheets: dict[str, Worksheet] = {}
    for unit_label in kosei.unit_labels:
        sheet_title = kosei.unit_sheet_title(unit_label)
        if sheet_title not in sheets:
            raise ValueError(
                f"構成のユニット「{unit_label}」に対応するシート「{sheet_title}」がありません。"
                f"（既存シート: {', '.join(sheets.keys()) or 'なし'}）"
            )
        unit_sheets[unit_label] = sheets[sheet_title]

    mapped_titles = {kosei.unit_sheet_title(u) for u in kosei.unit_labels}
    extra = set(sheets.keys()) - mapped_titles
    if extra:
        raise ValueError(
            f"構成に無いユニットシートがあります: {', '.join(sorted(extra))}"
        )

    return unit_sheets
