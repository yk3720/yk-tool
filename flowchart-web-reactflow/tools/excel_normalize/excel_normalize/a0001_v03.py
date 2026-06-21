"""A0001 塗布装置 — v0.3 形式（10 ユニット × 10 モジュール）の Excel 生成。"""

from __future__ import annotations

from openpyxl import Workbook

from excel_normalize.constants import (
    KOSEI_HEADERS_V03,
    KOSEI_SHEET,
    TABLE_GAP_COLS,
)
from excel_normalize.workbook_builder import _add_flow_table

INTERNAL_CODE = "A0001"
DISPLAY_NAME = "塗布装置"
UNIT_COUNT = 10
MODULES_PER_UNIT = 10
ROW_COUNT = UNIT_COUNT * MODULES_PER_UNIT


def unit_label(uin_id: int) -> str:
    return f"ﾕﾆｯﾄ{uin_id}"


def module_label(uin_id: int, index: int) -> str:
    return f"動作{uin_id * 100 + index:03d}"


def module_id(uin_id: int, index: int) -> int:
    return uin_id * 100 + index


def flow_rows_for_module(module_name: str) -> list[list[str]]:
    middle = module_name if module_name != "動作000" else "ワーク取出"
    return [
        ["10", "端子", "20", "", "1", "0", "開始", "", "", ""],
        ["20", "処理", "30", "", "2", "0", middle, "", "", ""],
        ["30", "端子", "", "", "3", "0", "終了", "", "", ""],
    ]


def kosei_row_formulas(excel_row: int) -> tuple[str, str, str, str, str, str]:
    """構成シート1行分の数式（マスター表を XLOOKUP）。"""
    return (
        '=XLOOKUP("*",装置名!$A$2:$A$2,装置名!$A$2:$A$2,,2)',
        '=XLOOKUP("*",装置名!$A$2:$A$2,装置名!$B$2:$B$2,,2)',
        "=QUOTIENT(ROW()-2,10)",
        f"=XLOOKUP(C{excel_row},ユニット!$A:$A,ユニット!$B:$B)",
        f"=C{excel_row}*100+MOD(ROW()-2,10)",
        f"=XLOOKUP(E{excel_row},モジュール!$A:$A,モジュール!$B:$B)",
    )


def _populate_master_sheets(wb: Workbook) -> None:
    ws_device = wb.create_sheet("装置名")
    ws_device.append(["装置製番", "装置名"])
    ws_device.append([INTERNAL_CODE, DISPLAY_NAME])

    ws_units = wb.create_sheet("ユニット")
    ws_units.append(["ID", "ユニット名"])
    for uin_id in range(UNIT_COUNT):
        ws_units.append([uin_id, unit_label(uin_id)])

    ws_modules = wb.create_sheet("モジュール")
    ws_modules.append(["ID", "動作"])
    for uin_id in range(UNIT_COUNT):
        for mod_index in range(MODULES_PER_UNIT):
            ws_modules.append(
                [module_id(uin_id, mod_index), module_label(uin_id, mod_index)]
            )


def _populate_kosei_sheet(wb: Workbook) -> None:
    ws_kosei = wb.create_sheet(KOSEI_SHEET, 0)
    ws_kosei.append(list(KOSEI_HEADERS_V03))
    for idx in range(ROW_COUNT):
        excel_row = idx + 2
        for col, formula in enumerate(kosei_row_formulas(excel_row), start=1):
            ws_kosei.cell(excel_row, col, formula)


def _populate_unit_flow_sheets(wb: Workbook) -> None:
    for uin_id in range(UNIT_COUNT):
        ws_unit = wb.create_sheet(f"U{uin_id}")
        start_col = 1
        for mod_index in range(MODULES_PER_UNIT):
            mod_name = module_label(uin_id, mod_index)
            _add_flow_table(
                ws_unit,
                table_name=mod_name,
                start_col=start_col,
                start_row=1,
                data_rows=flow_rows_for_module(mod_name),
            )
            start_col += TABLE_GAP_COLS


def build_a0001_v03_workbook() -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)
    _populate_master_sheets(wb)
    _populate_kosei_sheet(wb)
    _populate_unit_flow_sheets(wb)
    return wb
